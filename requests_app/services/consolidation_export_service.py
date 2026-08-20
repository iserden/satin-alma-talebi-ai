from __future__ import annotations

import csv
from decimal import Decimal
from io import BytesIO, StringIO

from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from requests_app.models import ConsolidationBatch


class ConsolidationExportError(ValueError):
    """Birleştirme çıktısı üretilemediğinde oluşur."""


def get_export_lines(batch: ConsolidationBatch):
    """Birleştirilmiş ürün satırlarını getirir."""

    lines = list(
        batch.lines
        .select_related("product")
        .order_by(
            "product__item_name",
            "product__sap_item_code",
        )
    )

    if not lines:
        raise ConsolidationExportError(
            "Bu birleştirmede dışa aktarılabilecek ürün bulunmuyor."
        )

    return lines


def get_source_forms(batch: ConsolidationBatch):
    """Birleştirmeye dahil edilen kaynak formları getirir."""

    return [
        source.purchase_request
        for source in (
            batch.sources
            .select_related("purchase_request")
            .order_by("purchase_request_id")
        )
    ]


def build_export_filename(
    batch: ConsolidationBatch,
    extension: str,
) -> str:
    """Güvenli dosya adı oluşturur."""

    safe_title = slugify(batch.title) or "talep-birlestirmesi"

    return (
        f"birlesim-{batch.pk}-"
        f"{safe_title}.{extension}"
    )


def decimal_to_json_number(
    value: Decimal | None,
):
    """
    Decimal değerini JSON içinde sayı olarak kullanılabilecek
    değere dönüştürür.
    """

    if value is None:
        return 0

    if value == value.to_integral_value():
        return int(value)

    return float(value)


def decimal_to_csv_text(
    value: Decimal | None,
) -> str:
    """
    CSV dosyasında Türkçe Excel uyumluluğu için
    ondalık ayıracı virgül yapar.
    """

    if value is None:
        value = Decimal("0")

    return format(
        value,
        ".2f",
    ).replace(
        ".",
        ",",
    )


def build_consolidation_json_payload(
    batch: ConsolidationBatch,
) -> dict:
    """Birleştirme sonucunun JSON veri yapısını oluşturur."""

    lines = get_export_lines(batch)
    source_forms = get_source_forms(batch)

    created_at = timezone.localtime(
        batch.created_at
    )

    return {
        "batch_id": batch.pk,
        "title": batch.title,
        "status": batch.status,
        "status_display": batch.get_status_display(),
        "created_at": created_at.isoformat(),
        "created_by": (
            batch.created_by.get_username()
            if batch.created_by
            else None
        ),
        "source_form_count": len(source_forms),
        "source_forms": [
            {
                "request_id": source_form.pk,
                "business_name": source_form.business_name,
                "authorized_person": (
                    source_form.authorized_person
                ),
                "form_number": source_form.form_number,
                "form_date": (
                    source_form.form_date.isoformat()
                    if source_form.form_date
                    else None
                ),
            }
            for source_form in source_forms
        ],
        "item_count": len(lines),
        "items": [
            {
                "item_code": line.product.sap_item_code,
                "item_name": line.product.item_name,
                "unit": line.unit,
                "week1": decimal_to_json_number(
                    line.week1
                ),
                "week2": decimal_to_json_number(
                    line.week2
                ),
                "week3": decimal_to_json_number(
                    line.week3
                ),
                "week4": decimal_to_json_number(
                    line.week4
                ),
                "week5": decimal_to_json_number(
                    line.week5
                ),
                "total_quantity": decimal_to_json_number(
                    line.total_quantity
                ),
                "source_request_count": (
                    line.source_request_count
                ),
                "source_line_count": (
                    line.source_line_count
                ),
            }
            for line in lines
        ],
    }


def build_consolidation_csv(
    batch: ConsolidationBatch,
) -> bytes:
    """Birleştirme sonucunu CSV dosyasına dönüştürür."""

    lines = get_export_lines(batch)

    output = StringIO(
        newline="",
    )

    writer = csv.writer(
        output,
        delimiter=";",
        lineterminator="\n",
    )

    writer.writerow(
        [
            "sap_item_code",
            "item_name",
            "unit",
            "week1",
            "week2",
            "week3",
            "week4",
            "week5",
            "total_quantity",
            "source_request_count",
            "source_line_count",
        ]
    )

    for line in lines:
        writer.writerow(
            [
                line.product.sap_item_code,
                line.product.item_name,
                line.unit,
                decimal_to_csv_text(line.week1),
                decimal_to_csv_text(line.week2),
                decimal_to_csv_text(line.week3),
                decimal_to_csv_text(line.week4),
                decimal_to_csv_text(line.week5),
                decimal_to_csv_text(
                    line.total_quantity
                ),
                line.source_request_count,
                line.source_line_count,
            ]
        )

    # UTF-8 BOM, Türkçe karakterlerin Excel'de
    # doğru görünmesini sağlar.
    return (
        "\ufeff" + output.getvalue()
    ).encode("utf-8")


def apply_header_style(
    worksheet,
    cell_range,
):
    """Excel başlık hücrelerine kurumsal görünüm uygular."""

    fill = PatternFill(
        fill_type="solid",
        fgColor="004070",
    )

    font = Font(
        color="FFFFFF",
        bold=True,
    )

    for row in worksheet[cell_range]:
        for cell in row:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )


def build_consolidation_excel(
    batch: ConsolidationBatch,
) -> BytesIO:
    """Birleştirme sonucunu biçimlendirilmiş Excel olarak üretir."""

    lines = get_export_lines(batch)
    source_forms = get_source_forms(batch)

    created_at = timezone.localtime(
        batch.created_at
    )

    workbook = Workbook()

    # =================================================
    # ÖZET SAYFASI
    # =================================================

    summary_sheet = workbook.active
    summary_sheet.title = "Ozet"
    summary_sheet.sheet_view.showGridLines = False

    summary_sheet.merge_cells(
        "A1:D1"
    )

    summary_sheet["A1"] = (
        "SATIN ALMA TALEP BİRLEŞTİRME ÖZETİ"
    )

    summary_sheet["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="004070",
    )

    summary_sheet["A1"].font = Font(
        color="FFFFFF",
        bold=True,
        size=14,
    )

    summary_sheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    summary_sheet.row_dimensions[1].height = 28

    summary_rows = [
        ["Birleştirme ID", batch.pk],
        ["Birleştirme adı", batch.title],
        ["Durum", batch.get_status_display()],
        [
            "Oluşturulma zamanı",
            created_at.strftime("%d.%m.%Y %H:%M"),
        ],
        [
            "Oluşturan kullanıcı",
            (
                batch.created_by.get_username()
                if batch.created_by
                else "Sistem"
            ),
        ],
        ["Kaynak form sayısı", len(source_forms)],
        ["Farklı ürün sayısı", len(lines)],
    ]

    for row_number, row_values in enumerate(
        summary_rows,
        start=3,
    ):
        summary_sheet.cell(
            row=row_number,
            column=1,
            value=row_values[0],
        )

        summary_sheet.cell(
            row=row_number,
            column=2,
            value=row_values[1],
        )

        summary_sheet.cell(
            row=row_number,
            column=1,
        ).font = Font(
            bold=True,
            color="003253",
        )

        summary_sheet.cell(
            row=row_number,
            column=1,
        ).fill = PatternFill(
            fill_type="solid",
            fgColor="EAF3F8",
        )

    summary_sheet.column_dimensions["A"].width = 28
    summary_sheet.column_dimensions["B"].width = 48
    summary_sheet.column_dimensions["C"].width = 18
    summary_sheet.column_dimensions["D"].width = 18

    # =================================================
    # BİRLEŞTİRİLMİŞ ÜRÜNLER SAYFASI
    # =================================================

    items_sheet = workbook.create_sheet(
        title="Birlesmis_Urunler"
    )

    items_sheet.sheet_view.showGridLines = False

    headers = [
        "SAP Kodu",
        "Ürün Adı",
        "Birim",
        "1. Hafta",
        "2. Hafta",
        "3. Hafta",
        "4. Hafta",
        "5. Hafta",
        "Toplam",
        "Kaynak Form",
        "Kaynak Satır",
    ]

    items_sheet.append(headers)

    for line in lines:
        items_sheet.append(
            [
                line.product.sap_item_code,
                line.product.item_name,
                line.unit,
                line.week1,
                line.week2,
                line.week3,
                line.week4,
                line.week5,
                line.total_quantity,
                line.source_request_count,
                line.source_line_count,
            ]
        )

    apply_header_style(
        items_sheet,
        f"A1:K1",
    )

    items_sheet.freeze_panes = "A2"
    items_sheet.auto_filter.ref = (
        f"A1:K{items_sheet.max_row}"
    )

    column_widths = {
        "A": 18,
        "B": 32,
        "C": 13,
        "D": 14,
        "E": 14,
        "F": 14,
        "G": 14,
        "H": 14,
        "I": 15,
        "J": 15,
        "K": 15,
    }

    for column_letter, width in column_widths.items():
        items_sheet.column_dimensions[
            column_letter
        ].width = width

    for row_number in range(
        2,
        items_sheet.max_row + 1,
    ):
        for column_number in range(
            4,
            10,
        ):
            items_sheet.cell(
                row=row_number,
                column=column_number,
            ).number_format = "#,##0.00"

        items_sheet.cell(
            row=row_number,
            column=1,
        ).number_format = "@"

    # =================================================
    # KAYNAK FORMLAR SAYFASI
    # =================================================

    source_sheet = workbook.create_sheet(
        title="Kaynak_Formlar"
    )

    source_sheet.sheet_view.showGridLines = False

    source_headers = [
        "Talep ID",
        "İşletme",
        "Yetkili Personel",
        "Form Numarası",
        "Form Tarihi",
        "Durum",
    ]

    source_sheet.append(source_headers)

    for source_form in source_forms:
        source_sheet.append(
            [
                source_form.pk,
                source_form.business_name,
                source_form.authorized_person,
                source_form.form_number,
                source_form.form_date,
                source_form.get_status_display(),
            ]
        )

    apply_header_style(
        source_sheet,
        "A1:F1",
    )

    source_sheet.freeze_panes = "A2"
    source_sheet.auto_filter.ref = (
        f"A1:F{source_sheet.max_row}"
    )

    source_widths = {
        "A": 12,
        "B": 30,
        "C": 28,
        "D": 22,
        "E": 16,
        "F": 18,
    }

    for column_letter, width in source_widths.items():
        source_sheet.column_dimensions[
            column_letter
        ].width = width

    for row_number in range(
        2,
        source_sheet.max_row + 1,
    ):
        source_sheet.cell(
            row=row_number,
            column=5,
        ).number_format = "DD.MM.YYYY"

    output = BytesIO()

    workbook.save(output)
    output.seek(0)

    return output


def mark_batch_as_exported(
    batch: ConsolidationBatch,
) -> None:
    """
    İlk dışa aktarmadan sonra durumu EXPORTED yapar.

    SAP'ye aktarılmış kayıt tekrar EXPORTED durumuna düşürülmez.
    """

    if batch.status != ConsolidationBatch.Status.READY:
        return

    ConsolidationBatch.objects.filter(
        pk=batch.pk,
        status=ConsolidationBatch.Status.READY,
    ).update(
        status=ConsolidationBatch.Status.EXPORTED,
    )

    batch.status = ConsolidationBatch.Status.EXPORTED