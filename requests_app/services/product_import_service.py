from __future__ import annotations

import re
from typing import Any

from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from requests_app.models import Product, ProductAlias
from requests_app.utils import normalize_product_text


EXPECTED_HEADERS = [
    "sap_item_code",
    "item_name",
    "unit",
    "category",
    "aliases",
    "is_active",
]

REQUIRED_HEADERS = {
    "sap_item_code",
    "item_name",
    "unit",
}

MAXIMUM_DATA_ROWS = 10_000


class ProductImportError(ValueError):
    """Ürün kataloğu içe aktarma hatalarını temsil eder."""


def cell_to_text(value: Any) -> str:
    """
    Excel hücresini güvenli metne dönüştürür.

    Excel'de tam sayı olarak saklanan 101.0 gibi değerlerin
    101 olarak okunmasını sağlar.
    """

    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value)).strip()

    return str(value).strip()


def parse_is_active(value: Any, row_number: int) -> bool:
    """Excel'deki aktif/pasif değerini Boolean değere dönüştürür."""

    normalized_value = normalize_product_text(
        cell_to_text(value)
    )

    if not normalized_value:
        return True

    true_values = {
        "true",
        "1",
        "yes",
        "evet",
        "aktif",
        "active",
    }

    false_values = {
        "false",
        "0",
        "no",
        "hayir",
        "pasif",
        "inactive",
    }

    if normalized_value in true_values:
        return True

    if normalized_value in false_values:
        return False

    raise ProductImportError(
        f"{row_number}. satırdaki is_active değeri geçersiz: "
        f"{cell_to_text(value)}"
    )


def split_aliases(value: Any) -> list[str]:
    """
    Alternatif isimleri ayırır.

    Desteklenen ayraçlar:
    - Dikey çizgi: |
    - Alt satır
    """

    text = cell_to_text(value)

    if not text:
        return []

    aliases = re.split(
        pattern=r"[|\r\n]+",
        string=text,
    )

    unique_aliases: list[str] = []
    used_normalized_aliases: set[str] = set()

    for alias in aliases:
        cleaned_alias = alias.strip()
        normalized_alias = normalize_product_text(cleaned_alias)

        if not normalized_alias:
            continue

        if normalized_alias in used_normalized_aliases:
            continue

        used_normalized_aliases.add(normalized_alias)
        unique_aliases.append(cleaned_alias)

    return unique_aliases


def build_header_map(header_row: tuple[Any, ...]) -> dict[str, int]:
    """Başlık isimlerini Excel sütun indeksleriyle eşleştirir."""

    header_map: dict[str, int] = {}
    duplicate_headers: set[str] = set()

    for column_index, raw_header in enumerate(header_row):
        header_name = cell_to_text(raw_header).casefold()

        if not header_name:
            continue

        if header_name in header_map:
            duplicate_headers.add(header_name)
            continue

        header_map[header_name] = column_index

    if duplicate_headers:
        raise ProductImportError(
            "Excel dosyasında tekrar eden sütun başlıkları var: "
            + ", ".join(sorted(duplicate_headers))
        )

    missing_headers = REQUIRED_HEADERS - set(header_map)

    if missing_headers:
        raise ProductImportError(
            "Excel dosyasında zorunlu sütunlar eksik: "
            + ", ".join(sorted(missing_headers))
        )

    return header_map


def get_row_value(
    row: tuple[Any, ...],
    header_map: dict[str, int],
    field_name: str,
) -> Any:
    """İlgili başlığa karşılık gelen hücreyi döndürür."""

    column_index = header_map.get(field_name)

    if column_index is None:
        return None

    if column_index >= len(row):
        return None

    return row[column_index]


def parse_excel_rows(excel_file) -> tuple[list[dict[str, Any]], int]:
    """
    Excel dosyasını doğrular ve veritabanına hazır kayıtlar üretir.

    Dönen değer:
    - Hazırlanan ürün kayıtları
    - Atlanan boş satır sayısı
    """

    try:
        excel_file.seek(0)

        workbook = load_workbook(
            filename=excel_file,
            read_only=True,
            data_only=True,
        )

    except (
        InvalidFileException,
        OSError,
        ValueError,
    ) as exc:
        raise ProductImportError(
            "Excel dosyası açılamadı. "
            "Dosyanın geçerli bir .xlsx dosyası olduğundan emin olun."
        ) from exc

    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)

        header_row = next(rows, None)

        if header_row is None:
            raise ProductImportError(
                "Excel dosyası tamamen boş."
            )

        header_map = build_header_map(header_row)

        parsed_records: list[dict[str, Any]] = []
        skipped_rows = 0
        validation_errors: list[str] = []
        used_sap_codes: set[str] = set()

        for excel_row_number, row in enumerate(
            rows,
            start=2,
        ):
            if excel_row_number > MAXIMUM_DATA_ROWS + 1:
                raise ProductImportError(
                    f"Excel dosyasında en fazla "
                    f"{MAXIMUM_DATA_ROWS} ürün satırı bulunabilir."
                )

            if all(
                value is None or cell_to_text(value) == ""
                for value in row
            ):
                skipped_rows += 1
                continue

            sap_item_code = cell_to_text(
                get_row_value(
                    row,
                    header_map,
                    "sap_item_code",
                )
            )

            item_name = cell_to_text(
                get_row_value(
                    row,
                    header_map,
                    "item_name",
                )
            )

            unit = cell_to_text(
                get_row_value(
                    row,
                    header_map,
                    "unit",
                )
            ).upper()

            category = cell_to_text(
                get_row_value(
                    row,
                    header_map,
                    "category",
                )
            )

            aliases = split_aliases(
                get_row_value(
                    row,
                    header_map,
                    "aliases",
                )
            )

            if not sap_item_code:
                validation_errors.append(
                    f"{excel_row_number}. satır: "
                    "sap_item_code boş bırakılamaz."
                )

            if not item_name:
                validation_errors.append(
                    f"{excel_row_number}. satır: "
                    "item_name boş bırakılamaz."
                )

            if not unit:
                validation_errors.append(
                    f"{excel_row_number}. satır: "
                    "unit boş bırakılamaz."
                )

            normalized_sap_code = sap_item_code.casefold()

            if (
                normalized_sap_code
                and normalized_sap_code in used_sap_codes
            ):
                validation_errors.append(
                    f"{excel_row_number}. satır: "
                    f"{sap_item_code} SAP kodu Excel dosyasında "
                    "birden fazla kez kullanılmış."
                )

            if normalized_sap_code:
                used_sap_codes.add(normalized_sap_code)

            try:
                is_active = parse_is_active(
                    get_row_value(
                        row,
                        header_map,
                        "is_active",
                    ),
                    row_number=excel_row_number,
                )

            except ProductImportError as exc:
                validation_errors.append(str(exc))
                is_active = True

            parsed_records.append(
                {
                    "excel_row_number": excel_row_number,
                    "sap_item_code": sap_item_code,
                    "item_name": item_name,
                    "unit": unit,
                    "category": category,
                    "aliases": aliases,
                    "is_active": is_active,
                }
            )

        if validation_errors:
            visible_errors = validation_errors[:20]
            message = "\n".join(visible_errors)

            remaining_error_count = (
                len(validation_errors) - len(visible_errors)
            )

            if remaining_error_count > 0:
                message += (
                    f"\n... ve {remaining_error_count} hata daha var."
                )

            raise ProductImportError(message)

        if not parsed_records:
            raise ProductImportError(
                "Excel dosyasında aktarılabilecek ürün bulunamadı."
            )

        return parsed_records, skipped_rows

    finally:
        workbook.close()


def save_product_records(
    records: list[dict[str, Any]],
) -> dict[str, int]:
    """Doğrulanmış kayıtları ürün kataloğuna aktarır."""

    created_products = 0
    updated_products = 0
    created_aliases = 0
    existing_aliases = 0

    with transaction.atomic():
        for record in records:
            product, was_created = Product.objects.update_or_create(
                sap_item_code=record["sap_item_code"],
                defaults={
                    "item_name": record["item_name"],
                    "unit": record["unit"],
                    "category": record["category"],
                    "is_active": record["is_active"],
                },
            )

            if was_created:
                created_products += 1
            else:
                updated_products += 1

            for alias_name in record["aliases"]:
                normalized_alias = normalize_product_text(
                    alias_name
                )

                if not normalized_alias:
                    continue

                # Resmî ürün adıyla tamamen aynı alias'ı kaydetmeyiz.
                if normalized_alias == product.normalized_name:
                    continue

                existing_alias = ProductAlias.objects.filter(
                    product=product,
                    normalized_alias=normalized_alias,
                ).first()

                if existing_alias:
                    existing_aliases += 1

                    fields_to_update: list[str] = []

                    if not existing_alias.is_active:
                        existing_alias.is_active = True
                        fields_to_update.append("is_active")

                    if existing_alias.alias_name != alias_name:
                        existing_alias.alias_name = alias_name
                        fields_to_update.extend(
                            [
                                "alias_name",
                                "normalized_alias",
                            ]
                        )

                    if fields_to_update:
                        existing_alias.save(
                            update_fields=list(
                                dict.fromkeys(fields_to_update)
                            )
                        )

                    continue

                ProductAlias.objects.create(
                    product=product,
                    alias_name=alias_name,
                    is_active=True,
                )

                created_aliases += 1

    return {
        "created_products": created_products,
        "updated_products": updated_products,
        "created_aliases": created_aliases,
        "existing_aliases": existing_aliases,
    }


def import_product_catalog(excel_file) -> dict[str, int]:
    """Excel dosyasını doğrular ve katalog verilerini içe aktarır."""

    records, skipped_rows = parse_excel_rows(
        excel_file=excel_file,
    )

    summary = save_product_records(
        records=records,
    )

    summary["processed_rows"] = len(records)
    summary["skipped_rows"] = skipped_rows

    return summary