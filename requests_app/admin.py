from io import BytesIO

from django.contrib import admin, messages
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import path
from django.utils.html import format_html
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .forms import ProductExcelImportForm
from .models import (
    ConsolidationBatch,
    ConsolidationLine,
    ConsolidationSource,
    Product,
    ProductAlias,
    PurchaseRequest,
    PurchaseRequestLine,
)
from .services.product_import_service import (
    ProductImportError,
    import_product_catalog,
)


class ProductAliasInline(admin.TabularInline):
    model = ProductAlias
    extra = 1

    fields = [
        "alias_name",
        "is_active",
    ]


@admin.register(Product)
class ProductAdmin(admin.ModelAdmin):
    change_list_template = (
        "admin/requests_app/product/change_list.html"
    )

    import_excel_template = (
        "admin/requests_app/product/import_excel.html"
    )

    list_display = [
        "sap_item_code",
        "item_name",
        "unit",
        "category",
        "is_active",
    ]

    search_fields = [
        "sap_item_code",
        "item_name",
        "normalized_name",
        "aliases__alias_name",
    ]

    list_filter = [
        "category",
        "unit",
        "is_active",
    ]

    ordering = [
        "item_name",
    ]

    inlines = [
        ProductAliasInline,
    ]

    def get_urls(self):
        default_urls = super().get_urls()

        custom_urls = [
            path(
                "import-excel/",
                self.admin_site.admin_view(
                    self.import_excel_view
                ),
                name="requests_app_product_import_excel",
            ),
            path(
                "download-excel-template/",
                self.admin_site.admin_view(
                    self.download_excel_template_view
                ),
                name=(
                    "requests_app_product_download_excel_template"
                ),
            ),
        ]

        # Özel yolları önce yazmak önemlidir.
        return custom_urls + default_urls

    def import_excel_view(self, request):
        """Ürün kataloğunu Excel dosyasından içe aktarır."""

        if not self.has_add_permission(request):
            raise PermissionDenied

        if request.method == "POST":
            form = ProductExcelImportForm(
                request.POST,
                request.FILES,
            )

            if form.is_valid():
                try:
                    summary = import_product_catalog(
                        excel_file=form.cleaned_data[
                            "excel_file"
                        ],
                    )

                except ProductImportError as exc:
                    form.add_error(
                        "excel_file",
                        str(exc),
                    )

                except Exception as exc:
                    form.add_error(
                        "excel_file",
                        (
                            "Aktarım sırasında beklenmeyen bir "
                            f"hata oluştu: {exc}"
                        ),
                    )

                else:
                    self.message_user(
                        request,
                        (
                            "Excel aktarımı tamamlandı. "
                            f"Yeni ürün: "
                            f"{summary['created_products']}, "
                            f"güncellenen ürün: "
                            f"{summary['updated_products']}, "
                            f"yeni alternatif ad: "
                            f"{summary['created_aliases']}, "
                            f"mevcut alternatif ad: "
                            f"{summary['existing_aliases']}, "
                            f"atlanan boş satır: "
                            f"{summary['skipped_rows']}."
                        ),
                        level=messages.SUCCESS,
                    )

                    return redirect(
                        "admin:requests_app_product_changelist"
                    )

        else:
            form = ProductExcelImportForm()

        context = {
            **self.admin_site.each_context(request),
            "title": "Excel’den ürün kataloğu aktar",
            "opts": self.model._meta,
            "form": form,
        }

        request.current_app = self.admin_site.name

        return render(
            request,
            self.import_excel_template,
            context,
        )

    def download_excel_template_view(self, request):
        """Doldurulabilir ürün kataloğu Excel şablonu üretir."""

        if not self.has_view_permission(request):
            raise PermissionDenied

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Urunler"

        headers = [
            "sap_item_code",
            "item_name",
            "unit",
            "category",
            "aliases",
            "is_active",
        ]

        worksheet.append(headers)

        sample_rows = [
            [
                "ITM000101",
                "Patates",
                "KG",
                "Sebze ve Meyve",
                (
                    "patates kg|taze patates|"
                    "yemeklik patates"
                ),
                True,
            ],
            [
                "ITM000204",
                "Coca Cola 1 LT",
                "KOLİ",
                "İçecek",
                (
                    "kola 1 litre|coca cola 1lt|"
                    "coca 1 l"
                ),
                True,
            ],
            [
                "ITM000302",
                "Toz Şeker 1 KG",
                "PAKET",
                "Temel Gıda",
                "şeker|toz şeker|1 kg şeker",
                True,
            ],
        ]

        for sample_row in sample_rows:
            worksheet.append(sample_row)

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="004070",
        )

        header_font = Font(
            color="FFFFFF",
            bold=True,
        )

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = "A1:F4"

        worksheet.column_dimensions["A"].width = 20
        worksheet.column_dimensions["B"].width = 32
        worksheet.column_dimensions["C"].width = 14
        worksheet.column_dimensions["D"].width = 24
        worksheet.column_dimensions["E"].width = 55
        worksheet.column_dimensions["F"].width = 14

        # SAP kodunun başındaki sıfırların korunması için metin biçimi.
        for cell in worksheet["A"]:
            cell.number_format = "@"

        instructions_sheet = workbook.create_sheet(
            title="Talimatlar"
        )

        instructions = [
            [
                "Alan",
                "Açıklama",
            ],
            [
                "sap_item_code",
                "Zorunlu ve benzersiz SAP kalem kodu.",
            ],
            [
                "item_name",
                "Ürünün resmî katalog adı.",
            ],
            [
                "unit",
                "KG, KOLİ, PAKET, ADET, LİTRE gibi birim.",
            ],
            [
                "category",
                "İsteğe bağlı ürün kategorisi.",
            ],
            [
                "aliases",
                (
                    "Alternatif ürün adlarını | işaretiyle "
                    "ayırın."
                ),
            ],
            [
                "is_active",
                (
                    "TRUE/FALSE, EVET/HAYIR, "
                    "1/0 veya AKTİF/PASİF."
                ),
            ],
        ]

        for instruction in instructions:
            instructions_sheet.append(instruction)

        for cell in instructions_sheet[1]:
            cell.fill = header_fill
            cell.font = header_font

        instructions_sheet.column_dimensions["A"].width = 22
        instructions_sheet.column_dimensions["B"].width = 70

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
        )

        response[
            "Content-Disposition"
        ] = (
            'attachment; filename="urun_katalogu_sablonu.xlsx"'
        )

        return response


@admin.register(ProductAlias)
class ProductAliasAdmin(admin.ModelAdmin):
    list_display = [
        "alias_name",
        "product",
        "is_active",
    ]

    search_fields = [
        "alias_name",
        "normalized_alias",
        "product__item_name",
        "product__sap_item_code",
    ]

    list_filter = [
        "is_active",
    ]


class PurchaseRequestLineInline(admin.TabularInline):
    model = PurchaseRequestLine
    extra = 0

    fields = [
        "row_number",
        "product_name",
        "unit",
        "matched_product",
        "matching_status",
    ]


@admin.register(PurchaseRequest)
class PurchaseRequestAdmin(admin.ModelAdmin):
    list_display = [
        "id",
        "business_name",
        "form_number",
        "form_date",
        "status_badge",
        "approved_by",
        "approved_at",
        "created_at",
    ]

    readonly_fields = [
        "approved_by",
        "approved_at",
    ]

    search_fields = [
        "business_name",
        "form_number",
        "authorized_person",
    ]

    list_filter = [
        "status",
        "form_date",
    ]

    inlines = [
        PurchaseRequestLineInline,
    ]

    @admin.display(
        description="Durum",
        ordering="status",
    )
    def status_badge(self, obj):
        style_map = {
            PurchaseRequest.Status.UPLOADED: "info",
            PurchaseRequest.Status.PROCESSING: "warning",
            PurchaseRequest.Status.REVIEW_REQUIRED: "warning",
            PurchaseRequest.Status.APPROVED: "success",
            PurchaseRequest.Status.FAILED: "danger",
        }

        badge_style = style_map.get(
            obj.status,
            "neutral",
        )

        return format_html(
            '<span class="admin-badge admin-badge--{}">{}</span>',
            badge_style,
            obj.get_status_display(),
        )


@admin.register(PurchaseRequestLine)
class PurchaseRequestLineAdmin(admin.ModelAdmin):
    list_display = [
        "id",
        "purchase_request",
        "row_number",
        "product_name",
        "unit",
        "matched_product",
        "matching_status",
        "matching_confidence",
    ]

    search_fields = [
        "product_name",
        "raw_product_name",
        "matched_product__item_name",
        "matched_product__sap_item_code",
    ]

    list_filter = [
        "status",
        "matching_status",
        "matching_method",
        "unit",
    ]


class ConsolidationSourceInline(admin.TabularInline):
    model = ConsolidationSource
    extra = 0
    readonly_fields = [
        "purchase_request",
        "added_at",
    ]


class ConsolidationLineInline(admin.TabularInline):
    model = ConsolidationLine
    extra = 0
    readonly_fields = [
        "product",
        "unit",
        "week1",
        "week2",
        "week3",
        "week4",
        "week5",
        "source_request_count",
        "source_line_count",
    ]


@admin.register(ConsolidationBatch)
class ConsolidationBatchAdmin(admin.ModelAdmin):
    list_display = [
        "id",
        "title",
        "status",
        "created_by",
        "created_at",
    ]

    search_fields = [
        "title",
    ]

    list_filter = [
        "status",
        "created_at",
    ]

    readonly_fields = [
        "created_by",
        "created_at",
        "updated_at",
    ]

    inlines = [
        ConsolidationSourceInline,
        ConsolidationLineInline,
    ]


@admin.register(ConsolidationLine)
class ConsolidationLineAdmin(admin.ModelAdmin):
    list_display = [
        "batch",
        "product",
        "unit",
        "week1",
        "week2",
        "week3",
        "week4",
        "week5",
    ]

    search_fields = [
        "product__sap_item_code",
        "product__item_name",
        "batch__title",
    ]


@admin.register(ConsolidationSource)
class ConsolidationSourceAdmin(admin.ModelAdmin):
    list_display = [
        "batch",
        "purchase_request",
        "added_at",
    ]


admin.site.site_header = "Satın Alma Otomasyonu"
admin.site.site_title = "Satın Alma Yönetimi"
admin.site.index_title = "Sistem Yönetim Paneli"
